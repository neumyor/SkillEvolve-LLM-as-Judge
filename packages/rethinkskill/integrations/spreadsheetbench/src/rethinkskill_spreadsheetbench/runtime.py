"""Offline execution helpers for SpreadsheetBench code generation."""

from __future__ import annotations

import ast
import io
import mimetypes
import os
import re
import signal
import site
import subprocess
import sys
import tempfile
import textwrap
from pathlib import Path

from rethinkskill.utils.serde import snapshot_regular_file

_PATH_ASSIGNMENT = re.compile(
    r"^\s*(INPUT_PATH|OUTPUT_PATH)\s*=\s*.+$",
    re.MULTILINE,
)
_FENCED_PYTHON = re.compile(
    r"```(?:python|py)\s*\n(?P<code>.*?)```",
    re.DOTALL | re.IGNORECASE,
)
_FORBIDDEN_IMPORT_ROOTS = {
    "asyncio",
    "ctypes",
    "ftplib",
    "http",
    "multiprocessing",
    "requests",
    "shutil",
    "socket",
    "subprocess",
    "telnetlib",
    "urllib",
    "webbrowser",
}
_FORBIDDEN_CALL_NAMES = {"__import__", "compile", "eval", "exec"}
_FORBIDDEN_ATTRIBUTES = {
    "connect",
    "execv",
    "execve",
    "execvp",
    "execvpe",
    "fork",
    "forkpty",
    "kill",
    "killpg",
    "popen",
    "posix_spawn",
    "spawn",
    "system",
    "urlopen",
}


_RUNNER_TEMPLATE = textwrap.dedent(
    """
    import os
    import sys
    import traceback

    INPUT_PATH = {input_path!r}
    OUTPUT_PATH = {output_path!r}
    _ALLOWED_READ_ROOTS = {allowed_read_roots!r}
    _ALLOWED_READ_FILES = {allowed_read_files!r}
    _ALLOWED_LIST_DIRS = {allowed_list_dirs!r}
    _ALLOWED_WRITE_ROOT = {allowed_write_root!r}

    def _within(path, root):
        try:
            return (
                os.path.commonpath(
                    [os.path.realpath(path), os.path.realpath(root)]
                )
                == os.path.realpath(root)
            )
        except (TypeError, ValueError):
            return False

    def _path(value):
        if not isinstance(value, (str, bytes, os.PathLike)):
            raise RuntimeError(
                "RETHINKSKILL_GENERATED_CODE_ISOLATION_BLOCK: "
                "non-path filesystem target"
            )
        return os.fsdecode(value)

    def _require_write_path(event, value):
        path = _path(value)
        if not _within(path, _ALLOWED_WRITE_ROOT):
            raise RuntimeError(
                "RETHINKSKILL_GENERATED_CODE_FILE_SCOPE_BLOCK: "
                + event
                + ":"
                + os.path.realpath(path)
            )

    def _require_read_path(event, value):
        path = _path(value)
        allowed = (
            os.path.realpath(path) in _ALLOWED_READ_FILES
            or _within(path, _ALLOWED_WRITE_ROOT)
            or any(_within(path, root) for root in _ALLOWED_READ_ROOTS)
        )
        if not allowed:
            raise RuntimeError(
                "RETHINKSKILL_GENERATED_CODE_FILE_SCOPE_BLOCK: "
                + event
                + ":"
                + os.path.realpath(path)
            )

    def _require_list_path(event, value):
        path = _path(value)
        allowed = (
            os.path.realpath(path) in _ALLOWED_LIST_DIRS
            or _within(path, _ALLOWED_WRITE_ROOT)
            or any(_within(path, root) for root in _ALLOWED_READ_ROOTS)
        )
        if not allowed:
            raise RuntimeError(
                "RETHINKSKILL_GENERATED_CODE_FILE_SCOPE_BLOCK: "
                + event
                + ":"
                + os.path.realpath(path)
            )

    def _audit(event, args):
        if (
            event.startswith("socket.")
            or event.startswith("shutil.")
            or event.startswith("os.exec")
            or event.startswith("os.spawn")
            or event
            in {{
                "subprocess.Popen",
                "os.system",
                "os.posix_spawn",
                "os.fork",
                "os.forkpty",
                "os.kill",
                "pty.spawn",
            }}
        ):
            raise RuntimeError(
                "RETHINKSKILL_GENERATED_CODE_ISOLATION_BLOCK: " + event
            )
        if event == "open" and args:
            path = _path(args[0])
            mode = str(args[1] if len(args) > 1 else "r")
            flags = args[2] if len(args) > 2 and isinstance(args[2], int) else 0
            writing = any(flag in mode for flag in "wax+") or bool(
                flags
                & (
                    os.O_WRONLY
                    | os.O_RDWR
                    | os.O_APPEND
                    | os.O_CREAT
                    | os.O_TRUNC
                )
            )
            if writing:
                _require_write_path(event, path)
            else:
                _require_read_path(event, path)
        if event in {{
            "os.remove",
            "os.rmdir",
            "os.mkdir",
            "os.chmod",
            "os.chown",
            "os.truncate",
            "os.utime",
            "os.setxattr",
            "os.removexattr",
            "os.mknod",
            "os.mkfifo",
        }} and args:
            _require_write_path(event, args[0])
        if event == "os.rename" and len(args) >= 2:
            _require_write_path(event + ":src", args[0])
            _require_write_path(event + ":dst", args[1])
        if event in {{"os.link", "os.symlink"}}:
            raise RuntimeError(
                "RETHINKSKILL_GENERATED_CODE_ISOLATION_BLOCK: " + event
            )
        if event == "os.chdir" and args:
            _require_read_path(event, args[0])
        if event in {{"os.listdir", "os.scandir"}} and args:
            _require_list_path(event, args[0])

    sys.addaudithook(_audit)
    try:
    {user_code}
    except Exception:
        traceback.print_exc()
        sys.exit(2)
    """
)


def extract_python_code(response: str) -> str:
    """Extract one complete Python program from a provider response."""

    match = _FENCED_PYTHON.search(response)
    candidate = match.group("code").strip() if match else response.strip()
    if not candidate:
        raise ValueError("spreadsheet model response contains no Python code")
    try:
        ast.parse(candidate)
    except SyntaxError as exc:
        raise ValueError("spreadsheet_python_syntax_invalid") from exc
    return candidate


def validate_generated_code_isolation(code: str) -> None:
    """Reject generated code that can escape an offline workbook transform."""

    try:
        tree = ast.parse(code)
    except SyntaxError as exc:
        raise ValueError("spreadsheet_python_syntax_invalid") from exc
    violations: list[str] = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for alias in node.names:
                root = alias.name.split(".", 1)[0]
                if root in _FORBIDDEN_IMPORT_ROOTS:
                    violations.append(f"import:{alias.name}")
        elif isinstance(node, ast.ImportFrom):
            root = str(node.module or "").split(".", 1)[0]
            if root in _FORBIDDEN_IMPORT_ROOTS:
                violations.append(f"import:{node.module}")
        elif isinstance(node, ast.Call):
            if isinstance(node.func, ast.Name) and node.func.id in _FORBIDDEN_CALL_NAMES:
                violations.append(f"call:{node.func.id}")
            if isinstance(node.func, ast.Attribute) and node.func.attr in _FORBIDDEN_ATTRIBUTES:
                violations.append(f"attribute:{node.func.attr}")
        elif isinstance(node, ast.Attribute) and node.attr.startswith("__"):
            violations.append(f"dunder:{node.attr}")
    if violations:
        raise ValueError("RETHINKSKILL_GENERATED_CODE_ISOLATION_BLOCK")


def _stream_evidence(value: str) -> dict[str, object]:
    payload = value.encode("utf-8")
    return {
        "present": bool(payload),
        "bytes": len(payload),
    }


def run_generated_code(
    code: str,
    *,
    input_path: Path,
    output_path: Path,
    timeout_seconds: int = 120,
) -> dict[str, object]:
    """Execute a workbook transform in a scoped child process."""

    resolved_input = input_path.resolve()
    resolved_output = output_path.resolve()
    write_root = resolved_output.parent
    write_root.mkdir(parents=True, exist_ok=True)
    cleaned = _PATH_ASSIGNMENT.sub("", code)
    try:
        validate_generated_code_isolation(cleaned)
    except ValueError:
        return {
            "ok": False,
            "failure": "generated_code_rejected",
            "returncode": None,
            "timed_out": False,
            "stdout_evidence": _stream_evidence(""),
            "stderr_evidence": _stream_evidence(""),
        }
    with tempfile.NamedTemporaryFile(
        "w",
        suffix=".py",
        prefix=".rethinkskill-runner-",
        dir=write_root,
        delete=False,
        encoding="utf-8",
    ) as handle:
        runner_path = Path(handle.name)
    library_roots = {
        str(Path(path).resolve())
        for path in (
            *site.getsitepackages(),
            *sys.path,
        )
        if path
        and Path(path).is_dir()
        and ("site-packages" in Path(path).parts or "dist-packages" in Path(path).parts)
    }
    import_search_directories = {
        str(Path(path).resolve()) for path in sys.path if path and Path(path).is_dir()
    }
    script = _RUNNER_TEMPLATE.format(
        input_path=str(resolved_input),
        output_path=str(resolved_output),
        allowed_read_roots=sorted(
            {
                str(write_root),
                str(Path(sys.prefix).resolve()),
                str(Path(sys.base_prefix).resolve()),
                *library_roots,
            }
        ),
        allowed_read_files=sorted(
            {
                str(runner_path.resolve()),
                str(resolved_input),
                *(
                    str(Path(path).resolve())
                    for path in mimetypes.knownfiles
                    if Path(path).is_file()
                ),
            }
        ),
        allowed_list_dirs=sorted(import_search_directories),
        allowed_write_root=str(write_root),
        user_code=textwrap.indent(cleaned, "    "),
    )
    runner_path.write_text(script, encoding="utf-8")
    child_env = {
        "LANG": os.environ.get("LANG", "C.UTF-8"),
        "LC_ALL": os.environ.get("LC_ALL", "C.UTF-8"),
        "PATH": str(Path(sys.executable).resolve().parent),
        "TMPDIR": str(write_root),
        "TEMP": str(write_root),
        "TMP": str(write_root),
        "PYTHONDONTWRITEBYTECODE": "1",
        "PYTHONNOUSERSITE": "1",
    }
    try:
        process = subprocess.Popen(
            [sys.executable, str(runner_path)],
            cwd=write_root,
            env=child_env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            start_new_session=True,
        )
    except Exception:
        runner_path.unlink(missing_ok=True)
        raise
    timed_out = False
    try:
        stdout, stderr = process.communicate(timeout=timeout_seconds)
    except subprocess.TimeoutExpired:
        timed_out = True
        os.killpg(process.pid, signal.SIGTERM)
        try:
            stdout, stderr = process.communicate(timeout=5)
        except subprocess.TimeoutExpired:
            os.killpg(process.pid, signal.SIGKILL)
            stdout, stderr = process.communicate()
    finally:
        runner_path.unlink(missing_ok=True)
    failure = ""
    if timed_out:
        failure = "generated_code_timeout"
    elif process.returncode != 0:
        failure = "generated_code_nonzero_exit"
    elif not resolved_output.is_file():
        failure = "generated_code_missing_output"
    return {
        "ok": not failure,
        "failure": failure,
        "returncode": process.returncode,
        "timed_out": timed_out,
        "stdout_evidence": _stream_evidence(stdout),
        "stderr_evidence": _stream_evidence(stderr),
    }


def preview_workbook(path: Path) -> str:
    """Return a bounded deterministic preview without exposing golden data."""

    return preview_workbook_bytes(snapshot_regular_file(path).payload)


def preview_workbook_bytes(payload: bytes) -> str:
    """Return a preview from the exact bytes bound into task provenance."""

    if type(payload) is not bytes:
        raise TypeError("workbook preview payload must be exact bytes")

    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "SpreadsheetBench requires the optional rethinkskill-spreadsheetbench package"
        ) from exc
    workbook = openpyxl.load_workbook(
        io.BytesIO(payload),
        read_only=True,
        data_only=False,
    )
    lines: list[str] = []
    try:
        for sheet_name in workbook.sheetnames[:8]:
            sheet = workbook[sheet_name]
            lines.append(f"## Sheet: {sheet_name}")
            for row in sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row, 30),
                max_col=min(sheet.max_column, 16),
                values_only=True,
            ):
                values = ["" if value is None else str(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    lines.append("\t".join(values))
            if sheet.max_row > 30 or sheet.max_column > 16:
                lines.append(f"[preview truncated; dimensions={sheet.max_row}x{sheet.max_column}]")
    finally:
        workbook.close()
    return "\n".join(lines)[:30_000] or "(workbook has no visible cells)"
