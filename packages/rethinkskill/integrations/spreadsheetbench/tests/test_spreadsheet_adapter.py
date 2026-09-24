from __future__ import annotations

import json
import tempfile
import unittest
from collections.abc import Mapping
from pathlib import Path
from unittest.mock import patch

from rethinkskill_spreadsheetbench.plugin import spec
from rethinkskill_spreadsheetbench.runtime import (
    extract_python_code,
    run_generated_code,
    validate_generated_code_isolation,
)

from rethinkskill.benchmarks.capabilities import CapabilityCatalog
from rethinkskill.benchmarks.core import benchmark_catalog
from rethinkskill.benchmarks.harness_catalog import HarnessCatalog
from rethinkskill.benchmarks.scoring import Verdict, verify_spreadsheet
from rethinkskill.integrations.catalog import integration_catalog
from rethinkskill.runtime.runner import (
    NativeRunOptions,
    execute_native_plan,
    plan_native_run,
)
from rethinkskill.runtime.tasks import RenderedTask
from rethinkskill.runtime.types import ModelOutcome
from rethinkskill.utils.fs import Repository

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency-gated test module
    openpyxl = None


class SpreadsheetRuntimeSafetyTests(unittest.TestCase):
    def test_syntax_failure_does_not_persist_source_excerpt(self) -> None:
        secret = "spreadsheet-syntax-secret"
        with self.assertRaises(ValueError) as caught:
            extract_python_code(f"value = {secret!r}\nthis is not valid Python !")
        self.assertEqual(
            str(caught.exception),
            "spreadsheet_python_syntax_invalid",
        )
        self.assertNotIn(secret, str(caught.exception))

    def test_child_streams_are_summarized_without_persisting_raw_text(self) -> None:
        secret = "spreadsheet-child-secret"
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_path = root / "input.xlsx"
            input_path.write_bytes(b"fixture")
            result = run_generated_code(
                f"print({secret!r})\nraise RuntimeError({secret!r})",
                input_path=input_path,
                output_path=root / "output.xlsx",
            )
        serialized = json.dumps(result, sort_keys=True)
        self.assertFalse(result["ok"])
        self.assertEqual(result["failure"], "generated_code_nonzero_exit")
        self.assertNotIn(secret, serialized)
        self.assertNotIn("stdout", result)
        self.assertNotIn("stderr", result)
        self.assertTrue(result["stdout_evidence"]["present"])
        self.assertTrue(result["stderr_evidence"]["present"])
        self.assertNotIn("sha256", result["stdout_evidence"])


class _CodeExecutor:
    def __init__(self, code: str):
        self.code = code

    def public_manifest(self) -> Mapping[str, object]:
        return {"kind": "fake-codegen", "model_calls": 0}

    def execute(
        self,
        rendered: RenderedTask,
        *,
        workspace: Path,
        timeout_seconds: int,
    ) -> ModelOutcome:
        del rendered, workspace, timeout_seconds
        return ModelOutcome(
            status="COMPLETED",
            response=f"```python\n{self.code}\n```",
            raw="synthetic code generation",
            process={"returncode": 0, "timed_out": False},
            attempted_calls=1,
            completed_calls=1,
        )


@unittest.skipIf(openpyxl is None, "openpyxl is not installed")
class SpreadsheetNativeTests(unittest.TestCase):
    @staticmethod
    def _catalog() -> CapabilityCatalog:
        return CapabilityCatalog(
            benchmark_catalog(),
            HarnessCatalog((spec(),)),
            integration_catalog(),
        )

    def _repository(self, temporary: str) -> Repository:
        root = Path(temporary)
        (root / "pyproject.toml").write_text("", encoding="utf-8")
        return Repository.discover(root)

    def _workbook(
        self,
        path: Path,
        *,
        answer: float | None,
    ) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet["A1"] = 1
        sheet["A2"] = 2
        if answer is not None:
            sheet["B1"] = answer
        workbook.save(path)
        workbook.close()

    def test_native_spreadsheet_codegen_hides_golden_and_scores_artifact(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            repository = self._repository(temporary)
            dataset = repository.root / "dataset/test"
            dataset.mkdir(parents=True)
            dataset.joinpath("items.json").write_text(
                json.dumps(
                    [
                        {
                            "id": "sheet-1",
                            "instruction": "Put the sum of A1:A2 in B1.",
                            "instruction_type": "Cell-Level Manipulation",
                            "answer_position": "Sheet1!B1",
                            "spreadsheet_path": "spreadsheet/sheet-1",
                        }
                    ]
                ),
                encoding="utf-8",
            )
            corpus = repository.root / "corpus"
            task_directory = corpus / "spreadsheet/sheet-1"
            task_directory.mkdir(parents=True)
            self._workbook(
                task_directory / "1_sheet-1_init.xlsx",
                answer=None,
            )
            self._workbook(
                task_directory / "1_sheet-1_golden.xlsx",
                answer=3,
            )
            skill = repository.root / "skill.md"
            skill.write_text(
                "Use openpyxl and preserve the workbook.",
                encoding="utf-8",
            )
            plan = plan_native_run(
                repository=repository,
                catalog=self._catalog(),
                benchmark="spreadsheetbench",
                dataset=dataset.parent,
                skill=skill,
                output_root=repository.root / "runs/spreadsheet",
                options=NativeRunOptions(
                    split="test",
                    limit=1,
                    asset_root=corpus,
                ),
            )
            receipt = execute_native_plan(
                plan,
                executor=_CodeExecutor(
                    "from openpyxl import load_workbook\n"
                    "workbook = load_workbook(INPUT_PATH)\n"
                    "sheet = workbook['Sheet1']\n"
                    "sheet['B1'] = sheet['A1'].value + sheet['A2'].value\n"
                    "workbook.save(OUTPUT_PATH)\n"
                ),
                authorized=True,
            )
            task_root = repository.root / "runs/spreadsheet/tasks/sheet-1"
            row = json.loads(
                (repository.root / "runs/spreadsheet/results.jsonl").read_text(encoding="utf-8")
            )
            golden_in_workspace = (task_root / "workspace/cases/000/golden.xlsx").exists()
            golden_in_verifier = (task_root / "verifier/cases/000/golden.xlsx").is_file()
        self.assertFalse(golden_in_workspace)
        self.assertTrue(golden_in_verifier)
        self.assertEqual(row["hard"], 1, row)
        self.assertEqual(row["soft"], 1.0)
        self.assertEqual(row["metrics"]["case_fraction"], 1.0)
        self.assertEqual(
            receipt["status"],
            "RETHINKSKILL_NATIVE_EXECUTION_VALIDATED",
        )

    def test_official_numeric_quantization_is_not_strict_float_equality(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            prediction = root / "prediction.xlsx"
            reference = root / "reference.xlsx"
            self._workbook(prediction, answer=1.233)
            self._workbook(reference, answer=1.234)
            verification = verify_spreadsheet(
                prediction,
                reference,
                "Sheet1!B1",
            )
        self.assertIs(verification.verdict, Verdict.PASS)

    def test_generated_code_isolation_blocks_process_escape(self) -> None:
        with self.assertRaisesRegex(
            ValueError,
            "ISOLATION_BLOCK",
        ):
            validate_generated_code_isolation(
                "import subprocess\nsubprocess.run(['echo', 'escape'])"
            )
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_path = root / "input.xlsx"
            output_path = root / "output.xlsx"
            self._workbook(input_path, answer=None)
            result = run_generated_code(
                "import subprocess\nsubprocess.run(['echo', 'escape'])",
                input_path=input_path,
                output_path=output_path,
            )
        self.assertFalse(result["ok"])
        self.assertFalse(output_path.exists())

    def test_artifact_launch_failure_is_sanitized_and_accounted(self) -> None:
        secret = "spreadsheet-launch-secret"
        with tempfile.TemporaryDirectory() as temporary:
            repository = self._repository(temporary)
            dataset = repository.root / "dataset/test"
            dataset.mkdir(parents=True)
            dataset.joinpath("items.json").write_text(
                json.dumps(
                    [
                        {
                            "id": "sheet-1",
                            "instruction": "Copy A1 into B1.",
                            "instruction_type": "Cell-Level Manipulation",
                            "answer_position": "Sheet1!B1",
                            "spreadsheet_path": "spreadsheet/sheet-1",
                        }
                    ]
                ),
                encoding="utf-8",
            )
            corpus = repository.root / "corpus"
            task_directory = corpus / "spreadsheet/sheet-1"
            task_directory.mkdir(parents=True)
            self._workbook(task_directory / "1_sheet-1_init.xlsx", answer=None)
            self._workbook(task_directory / "1_sheet-1_golden.xlsx", answer=1)
            skill = repository.root / "skill.md"
            skill.write_text("Use openpyxl.", encoding="utf-8")
            plan = plan_native_run(
                repository=repository,
                catalog=self._catalog(),
                benchmark="spreadsheetbench",
                dataset=dataset.parent,
                skill=skill,
                output_root=repository.root / "runs/spreadsheet-failure",
                options=NativeRunOptions(
                    split="test",
                    limit=1,
                    asset_root=corpus,
                ),
            )
            with patch(
                "rethinkskill_spreadsheetbench.runtime.subprocess.Popen",
                side_effect=OSError(secret),
            ):
                receipt = execute_native_plan(
                    plan,
                    executor=_CodeExecutor(
                        "from openpyxl import load_workbook\n"
                        "workbook = load_workbook(INPUT_PATH)\n"
                        "workbook.save(OUTPUT_PATH)\n"
                    ),
                    authorized=True,
                )
            row = json.loads(
                (repository.root / "runs/spreadsheet-failure/results.jsonl").read_text(
                    encoding="utf-8"
                )
            )
        self.assertEqual(row["failure_class"], "artifact_runner_error")
        self.assertEqual(row["phase"], "artifact")
        self.assertEqual(row["attempted_calls"], 1)
        self.assertEqual(row["completed_calls"], 1)
        self.assertNotIn(secret, json.dumps(row))
        self.assertEqual(receipt["attempted_calls"], 1)
