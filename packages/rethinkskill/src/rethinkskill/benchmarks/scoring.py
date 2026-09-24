"""RethinkSkill benchmarks scoring."""

from __future__ import annotations

import ast
import datetime
import json
import math
import re
import string
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from enum import Enum
from pathlib import Path
from typing import Any

from rethinkskill.errors import ResultValidationError
from rethinkskill.utils.serde import freeze_json_mapping, strict_json_loads, thaw_json_mapping


class Verdict(str, Enum):
    PASS = "PASS"
    FAIL = "FAIL"
    ABSTAIN = "ABSTAIN"


@dataclass(frozen=True, slots=True)
class Verification:
    verdict: Verdict
    reason: str
    answer: str | None = None
    checked: int | None = None
    metrics: Mapping[str, float] | None = None

    def validate(self) -> None:
        """Reject malformed scorer/plugin results at the trust boundary."""
        if type(self.verdict) is not Verdict:
            raise ResultValidationError("verification verdict must be a Verdict value")
        if type(self.reason) is not str:
            raise ResultValidationError("verification reason must be a string")
        if self.answer is not None and type(self.answer) is not str:
            raise ResultValidationError("verification answer must be a string or None")
        if self.checked is not None and (type(self.checked) is not int or self.checked < 0):
            raise ResultValidationError(
                "verification checked must be a non-negative integer or None"
            )
        if self.metrics is not None:
            if not isinstance(self.metrics, Mapping):
                raise ResultValidationError("verification metrics must be a mapping or None")
            for name, value in self.metrics.items():
                if type(name) is not str or not name:
                    raise ResultValidationError(
                        "verification metric names must be non-empty strings"
                    )
                if type(value) not in (int, float) or not math.isfinite(value):
                    raise ResultValidationError(
                        f"verification metric values must be finite numbers: {name}"
                    )

    def to_dict(self) -> dict[str, object]:
        self.validate()
        value: dict[str, object] = {"verdict": self.verdict.value, "reason": self.reason}
        if self.answer is not None:
            value["answer"] = self.answer
        if self.checked is not None:
            value["checked"] = self.checked
        if self.metrics is not None:
            value["metrics"] = thaw_json_mapping(self.metrics)
        return value


def freeze_verification(value: Verification) -> Verification:
    """Detach one exact verifier result before core scoring or persistence."""
    if type(value) is not Verification:
        raise ResultValidationError("verifier must return an exact Verification")
    try:
        metrics = freeze_json_mapping(value.metrics) if value.metrics is not None else None
    except (TypeError, ValueError, UnicodeEncodeError) as exc:
        raise ResultValidationError("verification metrics must be finite canonical JSON") from exc
    frozen = Verification(
        verdict=value.verdict,
        reason=value.reason,
        answer=value.answer,
        checked=value.checked,
        metrics=metrics,
    )
    frozen.validate()
    return frozen


ANSWER_TAG = re.compile("<answer>(.*?)</answer>", re.IGNORECASE | re.DOTALL)


def extract_freeform(text: str) -> str:
    matches = ANSWER_TAG.findall(str(text or ""))
    if matches:
        return matches[-1].strip()
    lines = [line.strip() for line in str(text or "").splitlines() if line.strip()]
    return lines[-1] if lines else ""


def extract_strict_json(text: str) -> tuple[str, str | None]:
    raw = str(text or "").strip()
    try:
        value = strict_json_loads(raw)
    except json.JSONDecodeError:
        return ("", "invalid_json")
    if (
        not isinstance(value, dict)
        or set(value) != {"answer"}
        or (not isinstance(value["answer"], str))
    ):
        return ("", "schema_mismatch")
    return (value["answer"].strip(), None)


def decode_aliases(value: Any) -> list[str]:
    if isinstance(value, (list, tuple)):
        return [str(item) for item in value]
    text = str(value or "").strip()
    if not text:
        return [""]
    if text.startswith("["):
        try:
            parsed = strict_json_loads(text)
        except json.JSONDecodeError:
            parsed = ast.literal_eval(text)
        if isinstance(parsed, list):
            return [str(item) for item in parsed]
    return [text]


def parse_answer(response: str, *, strict_json: bool) -> tuple[str, str | None]:
    return extract_strict_json(response) if strict_json else (extract_freeform(response), None)


def verify_alfworld(*, won: bool | None = None) -> Verification:
    if won is True:
        return Verification(Verdict.PASS, "alfworld_environment_won", metrics={"success": 1.0})
    if won is False:
        return Verification(Verdict.FAIL, "alfworld_environment_not_won", metrics={"success": 0.0})
    return Verification(Verdict.ABSTAIN, "no_independent_goal_state_snapshot_verifier_frozen")


def _transform(value: object) -> object:
    if isinstance(value, bool):
        return round(float(value), 2)
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    if isinstance(value, datetime.time):
        return str(value)[:-3]
    if isinstance(value, datetime.datetime):
        origin = datetime.datetime(1899, 12, 30)
        delta = value - origin
        return round(delta.days + delta.seconds / 86400.0, 0)
    if isinstance(value, str):
        try:
            return round(float(value), 2)
        except ValueError:
            return value
    return value


def _values_match(observed: object, gold: object) -> bool:
    observed = _transform(observed)
    gold = _transform(gold)
    if observed == "" and gold is None or (observed is None and gold == ""):
        return True
    if type(observed) is not type(gold):
        return False
    return observed == gold


def verify_spreadsheet(prediction: Path, reference: Path, answer_position: str) -> Verification:
    try:
        import openpyxl
        from openpyxl.utils.cell import range_boundaries
    except ImportError:
        return Verification(Verdict.ABSTAIN, "openpyxl_dependency_unavailable")
    if not prediction.is_file() or not reference.is_file():
        return Verification(Verdict.FAIL, "workbook_missing")
    try:
        predicted = openpyxl.load_workbook(prediction, data_only=True, read_only=True)
        expected = openpyxl.load_workbook(reference, data_only=True, read_only=True)
    except Exception as exc:
        return Verification(Verdict.FAIL, f"workbook_parse_error:{type(exc).__name__}")
    checked = 0
    try:
        for raw in str(answer_position or "").split(","):
            part = raw.strip()
            if not part:
                continue
            if "!" in part:
                sheet, cell_range = part.split("!", 1)
                sheet = sheet.strip().strip("'\"")
            else:
                sheet, cell_range = (expected.sheetnames[0], part)
            cell_range = cell_range.strip().strip("'\"")
            if sheet not in predicted.sheetnames or sheet not in expected.sheetnames:
                return Verification(Verdict.FAIL, f"sheet_missing:{sheet}", checked=checked)
            min_col, min_row, max_col, max_row = range_boundaries(cell_range)
            for row in range(min_row, max_row + 1):
                for column in range(min_col, max_col + 1):
                    gold_value = expected[sheet].cell(row, column).value
                    observed_value = predicted[sheet].cell(row, column).value
                    checked += 1
                    if not _values_match(observed_value, gold_value):
                        return Verification(
                            Verdict.FAIL, f"value_mismatch:{sheet}!{row}:{column}", checked=checked
                        )
    finally:
        predicted.close()
        expected.close()
    if checked == 0:
        return Verification(Verdict.ABSTAIN, "answer_position_empty_or_unparseable", checked=0)
    return Verification(Verdict.PASS, "official_answer_cell_values_match", checked=checked)


def verify_livemath(
    response: str,
    choices: Sequence[Mapping[str, Any]],
    correct_choice: Mapping[str, Any],
    *,
    strict_json: bool = False,
) -> Verification:
    answer, parser_error = parse_answer(response, strict_json=strict_json)
    if parser_error:
        return Verification(Verdict.FAIL, parser_error, "")

    def normalize_label(value: object) -> str:
        return str(value).strip().upper().rstrip(".):")

    valid_labels = {normalize_label(choice.get("label", "")) for choice in choices}
    observed = normalize_label(answer)
    reason = "label"
    if observed not in valid_labels:
        observed_text = answer.lower()
        exact_text_labels = [
            normalize_label(choice.get("label", ""))
            for choice in choices
            if str(choice.get("text", "")).strip().lower() == observed_text
        ]
        if exact_text_labels:
            observed = exact_text_labels[0]
            reason = "choice_text"
        else:
            first_token = normalize_label(answer.split()[0]) if answer.split() else ""
            if first_token in valid_labels:
                observed = first_token
                reason = "first_token_label"
            else:
                reason = "unmatched_choice"
    correct_label = normalize_label(correct_choice.get("label", ""))
    exact = float(observed == correct_label)
    return Verification(
        Verdict.PASS if observed == correct_label else Verdict.FAIL,
        reason,
        observed or answer,
        metrics={"exact_match": exact, "f1": exact, "substring": exact},
    )


def normalized_searchqa(value: str) -> str:
    text = str(value).lower()
    text = "".join(character for character in text if character not in string.punctuation)
    text = re.sub("\\b(a|an|the)\\b", " ", text)
    return " ".join(text.split())


def normalized_officeqa(value: str) -> str:
    numeric_chars = set("0123456789.-")
    text = str(value or "").lower().strip().replace(",", "")
    text = "".join(
        character
        for character in text
        if character not in string.punctuation or character in numeric_chars or character == "%"
    )
    text = re.sub("\\b(?:million|millions|billion|billions|dollar|dollars|nominal)\\b", " ", text)
    return " ".join(text.split())


def verify_searchqa(response: str, aliases: Any) -> Verification:
    answer = extract_freeform(response)
    accepted = decode_aliases(aliases)
    normalized_answer = normalized_searchqa(answer)
    normalized_aliases = [normalized_searchqa(alias) for alias in accepted]
    exact = float(normalized_answer in normalized_aliases)
    substring = float(
        any(
            alias in normalized_answer or normalized_answer in alias for alias in normalized_aliases
        )
    )
    prediction_tokens = normalized_answer.split()
    best_f1 = 0.0
    for alias in normalized_aliases:
        gold_tokens = alias.split()
        if not prediction_tokens:
            best_f1 = max(best_f1, float(not gold_tokens))
            continue
        if not gold_tokens:
            continue
        common = Counter(prediction_tokens) & Counter(gold_tokens)
        overlap = sum(common.values())
        if overlap:
            precision = overlap / len(prediction_tokens)
            recall = overlap / len(gold_tokens)
            best_f1 = max(best_f1, 2 * precision * recall / (precision + recall))
    return Verification(
        verdict=Verdict.PASS if exact == 1.0 else Verdict.FAIL,
        reason="squad_exact_match",
        answer=answer,
        metrics={"exact_match": exact, "f1": best_f1, "substring": substring},
    )


def verify_officeqa(response: str, gold: str, *, strict_json: bool = False) -> Verification:
    answer, parser_error = parse_answer(response, strict_json=strict_json)
    if parser_error:
        return Verification(Verdict.FAIL, parser_error, "")
    predicted = normalized_officeqa(answer)
    expected = normalized_officeqa(gold)
    predicted_tokens = predicted.split()
    expected_tokens = expected.split()
    common = Counter(predicted_tokens) & Counter(expected_tokens)
    overlap = sum(common.values())
    f1 = 0.0
    if not predicted_tokens or not expected_tokens:
        f1 = float(predicted_tokens == expected_tokens)
    elif overlap:
        precision = overlap / len(predicted_tokens)
        recall = overlap / len(expected_tokens)
        f1 = 2 * precision * recall / (precision + recall)
    exact = float(predicted == expected)
    return Verification(
        Verdict.PASS if exact else Verdict.FAIL,
        "officeqa_normalized_exact",
        answer,
        metrics={"exact_match": exact, "f1": f1},
    )


def verify_docvqa(response: str, aliases: Any, *, strict_json: bool = False) -> Verification:
    answer, parser_error = parse_answer(response, strict_json=strict_json)
    if parser_error:
        return Verification(Verdict.FAIL, parser_error, "")
    predicted = " ".join(str(answer).strip().lower().split())
    score = 0.0
    for alias in decode_aliases(aliases):
        target = " ".join(str(alias).strip().lower().split())
        if predicted == target:
            score = 1.0
            break
        if not predicted or not target:
            continue
        previous = list(range(len(target) + 1))
        for index, predicted_character in enumerate(predicted, start=1):
            current = [index]
            for target_index, target_character in enumerate(target, start=1):
                current.append(
                    min(
                        current[target_index - 1] + 1,
                        previous[target_index] + 1,
                        previous[target_index - 1] + (predicted_character != target_character),
                    )
                )
            previous = current
        normalized_distance = previous[-1] / max(len(predicted), len(target))
        if normalized_distance < 0.5:
            score = max(score, 1.0 - normalized_distance)
    return Verification(
        Verdict.PASS if score >= 0.999 else Verdict.FAIL,
        "docvqa_anls",
        answer,
        metrics={"anls": score},
    )
